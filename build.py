#!/usr/bin/env python3
"""Gera o app do Caminho de Nhá Chica (index.html + sw.js) a partir da planilha
e dos PDFs das passagens.

Uso:
    python build.py                      # gera index.html e sw.js
    python build.py --fragmento saida.html   # também gera a versão sem <html>/<head>

Requer: openpyxl, pypdf, pillow.
"""
import argparse
import datetime as dt
import hashlib
import html
import json
import re
import sys
import urllib.parse
from pathlib import Path

import openpyxl
from PIL import ImageOps
from pypdf import PdfReader

RAIZ = Path(__file__).resolve().parent
PLANILHA = RAIZ / "Caminho_Nha_Chica_Roteiro.xlsx"
MODELO = RAIZ / "src" / "app.html"
SW_MODELO = RAIZ / "src" / "sw.js"

# Colunas da aba "Roteiro". C (dia da semana) é calculado aqui; K (Hospedagem
# Sugerida) fica de fora de propósito.
COL = {"etapa": "A", "data": "B", "dia": "D", "origem": "E", "destino": "F",
       "km": "G", "terreno": "H", "pouso": "I", "status": "J", "obs": "L"}

SEMANA = ["seg", "ter", "qua", "qui", "sex", "sáb", "dom"]
MESES = ["jan", "fev", "mar", "abr", "mai", "jun", "jul", "ago", "set", "out", "nov", "dez"]

# Dados copiados dos PDFs da Quero Passagem. O QR code é extraído do próprio PDF.
PASSAGENS = [
    {
        "id": "ida",
        "titulo": "Ida",
        "pdf": "IDA_20nov_Quero Passagem.pdf",
        "data": dt.date(2026, 11, 20),
        "partida": "15:30",
        "chegada": "19:30",
        "duracao": "4h",
        "origem": "São Paulo, SP",
        "origem_local": "Rodoviária do Tietê",
        "origem_mapa": "Terminal Rodoviário Tietê, São Paulo, SP",
        "destino": "Inconfidentes, MG",
        "destino_local": "",
        "empresa": "Santa Cruz · Executivo",
        "dados": [("Pedido", "RDSVTZ8E")],
        "qr_tipo": "app",
        "qr_legenda": "QR impresso no PDF. Ele abre a passagem no app Quero Passagem; escaneie de outro aparelho ou use o botão.",
        "aviso": "O PDF da ida não tem QR de embarque. Embarque mostrando a passagem no app Quero Passagem, junto com documento com foto. Deixe a passagem aberta no app antes de sair de casa.",
        "links": [
            ("Abrir passagem no app", "https://queropassagem.com.br/redirect-app?tipo=bpe&compra=6c8dc0e0286b45044cd8122106fbe47b"),
            ("Ver a compra no site", "https://queropassagem.com.br/compra-finalizada/6c8dc0e0286b45044cd8122106fbe47b"),
        ],
        "pagamento": [("1 assento", "R$ 169,39"), ("Seguro", "R$ 17,95"),
                      ("Taxa de serviço", "R$ 33,87"), ("Total (cartão de crédito)", "R$ 221,21")],
    },
    {
        "id": "volta",
        "titulo": "Volta",
        "pdf": "Volta_01dez_Quero Passagem.pdf",
        "data": dt.date(2026, 12, 1),
        "partida": "11:30",
        "chegada": "",
        "duracao": "",
        "origem": "Caxambu, MG",
        "origem_local": "Praça Cônego José Castilho Moreira, 10",
        "origem_mapa": "Praça Cônego José Castilho Moreira, 10, Caxambu, MG",
        "destino": "São Paulo, SP",
        "destino_local": "Rodoviária do Tietê · Av. Cruzeiro do Sul, 1800, Santana",
        "empresa": "Cometa · linha São Paulo (SP) – Cruzília (MG)",
        "dados": [("Passageiro", "Rodrigo Ramos Garcia"), ("Poltrona", "67")],
        "qr_tipo": "embarque",
        "qr_legenda": "QR de embarque. Mostre na porta do ônibus com o brilho da tela no máximo.",
        "aviso": "Leve documento com foto com os mesmos dados da reserva.",
        "links": [],
        "pagamento": [],
    },
]

SUPORTE = [("Quero Passagem (suporte)", "(11) 4680-2994")]
EMERGENCIA = [("SAMU", "192"), ("Bombeiros", "193"), ("Polícia Militar", "190")]

FONE = re.compile(r"\(\d{2}\)\s*\d[\d\s.-]{5,}\d")
MINUSCULAS = {"de", "da", "do", "das", "dos", "e"}


def esc(s):
    return html.escape(str(s), quote=True)


def km_txt(x):
    s = f"{x:.1f}".rstrip("0").rstrip(".")
    return s.replace(".", ",")


def quando(d):
    return f"{SEMANA[d.weekday()]}, {d.day} {MESES[d.month - 1]}"


def arrumar_nome(nome):
    """Converte nomes em CAIXA ALTA para Título; mantém os demais como estão."""
    if nome != nome.upper():
        return nome
    palavras = []
    for i, p in enumerate(nome.lower().split()):
        if i and p in MINUSCULAS:
            palavras.append(p)
        else:
            palavras.append(re.sub(r"[^\W\d_]", lambda m: m.group().upper(), p, count=1))
    return " ".join(palavras)


def separar_cidade(texto):
    """'Baependi (dorme em Caxambu)' -> ('Baependi', 'dorme em Caxambu')."""
    m = re.match(r"^(.*?)\s*\((.*)\)\s*$", texto or "")
    if not m:
        return (texto or "").strip(), ""
    nome, extra = m.group(1).strip(), m.group(2).strip()
    return nome, "" if extra == "MG" else extra


def tel_href(numero):
    return "tel:+55" + re.sub(r"\D", "", numero)


def ler_fones(linhas, avisos, onde):
    fones = []
    for linha in linhas:
        m = FONE.search(linha)
        if not m:
            continue
        numero = re.sub(r"\s+", " ", m.group().strip())
        quem = (linha[:m.start()] + linha[m.end():]).strip().strip("()").strip()
        digitos = re.sub(r"\D", "", numero)[2:]
        if digitos.startswith("9") and len(digitos) == 8:
            avisos.append(f"{onde}: {numero} tem 8 dígitos começando com 9 (celular deveria ter 9 dígitos)")
        fones.append({"numero": numero, "href": tel_href(numero), "quem": quem})
    return fones


def ler_roteiro(wb, avisos):
    ws = wb["Roteiro"]
    titulo = ws["A1"].value
    observacao = ""
    dias = []
    data = None
    for r in range(5, ws.max_row + 1):
        v = {k: ws[f"{c}{r}"].value for k, c in COL.items()}
        if isinstance(v["etapa"], str) and v["etapa"].startswith("Observa"):
            observacao = v["etapa"].split(":", 1)[1].strip()
            continue
        if not (isinstance(v["dia"], str) and v["dia"].startswith("Dia")):
            continue
        if isinstance(v["data"], dt.datetime):
            data = v["data"].date()
        else:
            m = re.fullmatch(r"=B\d+\+(\d+)", str(v["data"]))
            if not m or data is None:
                sys.exit(f"Linha {r}: data inesperada {v['data']!r}")
            data = data + dt.timedelta(days=int(m.group(1)))
        etapa = v["etapa"]
        if isinstance(etapa, float):
            etapa = f"Etapa {int(etapa)}"
        dias.append({**v, "linha": r, "data": data, "etapa": etapa,
                     "km": float(v["km"] or 0)})
    total = sum(d["km"] for d in dias)
    acum = 0.0
    for i, d in enumerate(dias):
        acum += d["km"]
        d["acum"] = acum
        d["faltam"] = total - acum
        d["id"] = "dia-" + d["dia"].split()[-1]
        d["origem_nome"], _ = separar_cidade(d["origem"])
        d["destino_nome"], d["destino_extra"] = separar_cidade(d["destino"])
        d["onibus"] = next((p for p in PASSAGENS if p["data"] == d["data"]), None)
        texto = (d["pouso"] or "").strip()
        if d["etapa"] == "Volta" or not texto:
            d["pouso"] = None
            continue
        linhas = [l.strip() for l in texto.splitlines() if l.strip()]
        status = [l.strip() for l in (d["status"] or "").splitlines() if l.strip()]
        m = re.search(r"dorme em (.+)", d["destino_extra"])
        cidade = m.group(1) if m else d["destino_nome"]
        nome = arrumar_nome(linhas[0])
        d["pouso"] = {
            "nome": nome,
            "fones": ler_fones(linhas[1:], avisos, f"{d['dia']} ({nome})"),
            "status": status[0] if status else "",
            "notas": status[1:],
            "cidade": cidade,
            "mapa": "https://www.google.com/maps/search/?api=1&query="
                    + urllib.parse.quote_plus(f"{nome}, {cidade}, MG"),
        }
    return titulo, dias, total, observacao


def ler_info(wb):
    ws = wb["Informações Gerais"]
    itens = []
    for r in range(3, ws.max_row + 1):
        k, v = ws[f"A{r}"].value, ws[f"B{r}"].value
        if k and v:
            itens.append((str(k).strip(), str(v).strip()))
    return itens


# ---------- QR code: lê os módulos da imagem do PDF e redesenha em SVG ----------

def qr_matriz(pdf):
    imagem = next(i.image for i in PdfReader(RAIZ / pdf).pages[0].images
                  if i.image.width == i.image.height)
    cinza = imagem.convert("L")
    x0, y0, x1, y1 = ImageOps.invert(cinza).point(lambda p: 255 if p > 127 else 0).getbbox()
    px = cinza.load()
    escuro = lambda x, y: px[int(x), int(y)] < 128
    corrida = 0
    while escuro(x0 + corrida, y0 + 1):
        corrida += 1
    n = round((x1 - x0) / (corrida / 7))
    n = min(range(21, 178, 4), key=lambda v: abs(v - n))
    passo = (x1 - x0) / n
    m = [[escuro(x0 + (c + .5) * passo, y0 + (r + .5) * passo) for c in range(n)] for r in range(n)]
    for r0, c0 in ((0, 0), (0, n - 7), (n - 7, 0)):
        borda = [m[r0][c0 + i] for i in range(7)] + [m[r0 + 6][c0 + i] for i in range(7)]
        miolo = [m[r0 + 1][c0 + i] for i in range(1, 6)]
        if not all(borda) or any(miolo):
            sys.exit(f"{pdf}: não reconheci o QR code")
    return m


def qr_svg(m, rotulo):
    q = 4
    lado = len(m) + 2 * q
    d = []
    for r, linha in enumerate(m):
        c = 0
        while c < len(linha):
            if linha[c]:
                ini = c
                while c < len(linha) and linha[c]:
                    c += 1
                d.append(f"M{ini + q} {r + q}h{c - ini}v1h-{c - ini}z")
            else:
                c += 1
    return (f'<svg class="qr-svg" viewBox="0 0 {lado} {lado}" role="img" aria-label="{esc(rotulo)}" '
            f'shape-rendering="crispEdges"><rect class="qr-fundo" width="{lado}" height="{lado}"/>'
            f'<path class="qr-tinta" d="{"".join(d)}"/></svg>')


# ---------- HTML ----------

ICONES = {
    "fone": '<svg class="ico" viewBox="0 0 24 24" aria-hidden="true"><path d="M6.6 10.8a15.2 15.2 0 0 0 6.6 6.6l2.2-2.2a1 1 0 0 1 1-.25 11.4 11.4 0 0 0 3.6.57 1 1 0 0 1 1 1V20a1 1 0 0 1-1 1A17 17 0 0 1 3 4a1 1 0 0 1 1-1h3.5a1 1 0 0 1 1 1c0 1.25.2 2.45.57 3.57a1 1 0 0 1-.25 1z"/></svg>',
    "mapa": '<svg class="ico" viewBox="0 0 24 24" aria-hidden="true"><path d="M12 2a7 7 0 0 0-7 7c0 5.25 7 13 7 13s7-7.75 7-13a7 7 0 0 0-7-7zm0 9.5A2.5 2.5 0 1 1 12 6.5a2.5 2.5 0 0 1 0 5z"/></svg>',
    "onibus": '<svg class="ico" viewBox="0 0 24 24" aria-hidden="true"><path d="M4 16c0 .88.39 1.67 1 2.22V20a1 1 0 0 0 1 1h1a1 1 0 0 0 1-1v-1h8v1a1 1 0 0 0 1 1h1a1 1 0 0 0 1-1v-1.78c.61-.55 1-1.34 1-2.22V6c0-3.5-3.58-4-8-4s-8 .5-8 4v10zm3.5 1a1.5 1.5 0 1 1 0-3 1.5 1.5 0 0 1 0 3zm9 0a1.5 1.5 0 1 1 0-3 1.5 1.5 0 0 1 0 3zM18 11H6V6h12v5z"/></svg>',
}


def html_fones(fones, classe="fones"):
    if not fones:
        return ""
    itens = "".join(
        f'<li><a class="fone num" href="{esc(f["href"])}">{ICONES["fone"]}{esc(f["numero"])}</a>'
        + (f' <span class="fone-quem">{esc(f["quem"])}</span>' if f["quem"] else "")
        + "</li>" for f in fones)
    return f'<ul class="{classe}">{itens}</ul>'


def html_selos(pouso):
    selos = []
    if pouso["status"]:
        cls = "ok" if pouso["status"].lower().startswith("confirm") else "neutro"
        selos.append(f'<span class="selo {cls}">{esc(pouso["status"])}</span>')
    for n in pouso["notas"]:
        cls = "alerta" if re.search(r"confirm|véspera", n, re.I) else "neutro"
        selos.append(f'<span class="selo {cls}">{esc(n)}</span>')
    return f'<p class="selos">{"".join(selos)}</p>' if selos else ""


def html_onibus(p):
    horario = p["partida"] + (f" – {p['chegada']}" if p["chegada"] else "")
    return (f'<a class="onibus-link" href="#{p["id"]}">{ICONES["onibus"]}'
            f'<span><strong>Ônibus {esc(p["origem"].split(",")[0])} → {esc(p["destino"].split(",")[0])}</strong>'
            f' <span class="num">{esc(horario)}</span><br><span class="onibus-ver">Ver passagem</span></span></a>')


def html_dia(d):
    caminhada = d["km"] > 0
    classes = "dia" + ("" if caminhada else " dia-parado")
    if d["origem_nome"] in ("", "—"):
        trecho = f'<span class="chegada-a">Chegada a</span> {esc(d["destino_nome"])}'
    else:
        trecho = f'{esc(d["origem_nome"])} <span class="seta" aria-label="até">→</span> {esc(d["destino_nome"])}'
    partes = [
        f'<li class="{classes}" id="{d["id"]}" data-data="{d["data"].isoformat()}">',
        f'<div class="marco" aria-hidden="true"><span class="marco-ponto"></span><span class="marco-km num">{km_txt(d["acum"])}</span></div>',
        '<div class="dia-corpo">',
        f'<p class="dia-quando"><strong>{esc(d["dia"])}</strong> · {quando(d["data"])} · {esc(d["etapa"])}'
        '<span class="hoje-tag">Hoje</span></p>',
        f'<h3 class="dia-trecho">{trecho}</h3>',
    ]
    if d["destino_extra"]:
        partes.append(f'<p class="dia-via">{esc(d["destino_extra"])}</p>')
    if caminhada:
        partes.append(
            f'<p class="dia-km"><span class="km-grande num">{km_txt(d["km"])}</span> km'
            f'<span class="dia-km-extra num">acumulado {km_txt(d["acum"])} · faltam {km_txt(d["faltam"])}</span></p>')
    if d["terreno"]:
        partes.append(f'<p class="dia-terreno">{esc(d["terreno"])}</p>')
    if d["onibus"]:
        partes.append(html_onibus(d["onibus"]))
    p = d["pouso"]
    if p:
        partes.append(
            '<div class="pouso">'
            '<p class="rotulo">Pouso</p>'
            f'<p class="pouso-nome">{esc(p["nome"])}</p>'
            + html_selos(p) + html_fones(p["fones"])
            + f'<a class="mapa-link" href="{esc(p["mapa"])}" target="_blank" rel="noopener">{ICONES["mapa"]}Procurar no mapa</a>'
            '</div>')
    if d["obs"] and d["obs"].strip() not in ("—", "-"):
        partes.append(f'<p class="dia-obs"><span class="rotulo">No caminho</span>{esc(d["obs"])}</p>')
    partes.append("</div></li>")
    return "".join(partes)


def html_faixa(dias):
    segs = "".join(
        f'<a class="seg" href="#{d["id"]}" style="flex-grow:{d["km"]}" data-dia="{d["id"]}" '
        f'aria-label="{esc(d["dia"])}: {km_txt(d["km"])} km"><span class="num">{km_txt(d["km"])}</span></a>'
        for d in dias if d["km"] > 0)
    andando = [d for d in dias if d["km"] > 0]
    return (
        '<nav class="faixa" aria-label="Etapas em escala de km; toque para ir ao dia">'
        f'<div class="faixa-barra">{segs}</div>'
        f'<div class="faixa-pontas"><span>{esc(andando[0]["origem_nome"])}</span>'
        f'<span class="faixa-meio num" id="faixa-meio">km por etapa</span><span>{esc(andando[-1]["destino_nome"])}</span></div>'
        '</nav>')


def html_passagem(p, qr):
    data_txt = f"{quando(p['data'])} {p['data'].year}"
    dados = "".join(f'<div><dt>{esc(k)}</dt><dd>{esc(v)}</dd></div>' for k, v in p["dados"])
    chegada = (f'<div class="b-ponto b-dir"><span class="b-hora num">{esc(p["chegada"])}</span>'
               if p["chegada"] else '<div class="b-ponto b-dir"><span class="b-hora" aria-hidden="true">&nbsp;</span>')
    chegada += f'<span class="b-local">{esc(p["destino"])}</span>'
    if p["destino_local"]:
        chegada += f'<span class="b-sub">{esc(p["destino_local"])}</span>'
    chegada += "</div>"
    mapa = ("https://www.google.com/maps/search/?api=1&query=" + urllib.parse.quote_plus(p["origem_mapa"]))
    links = "".join(
        f'<a class="botao{" botao-forte" if i == 0 else ""}" href="{esc(u)}" target="_blank" rel="noopener">{esc(t)}</a>'
        for i, (t, u) in enumerate(p["links"]))
    links += f'<a class="botao" href="{esc(mapa)}" target="_blank" rel="noopener">{ICONES["mapa"]}Local de embarque</a>'
    pagamento = ""
    if p["pagamento"]:
        linhas = "".join(f'<div><dt>{esc(k)}</dt><dd class="num">{esc(v)}</dd></div>' for k, v in p["pagamento"])
        pagamento = f'<details class="b-pag"><summary>Pagamento</summary><dl>{linhas}</dl></details>'
    qr_classe = "b-qr b-qr-embarque" if p["qr_tipo"] == "embarque" else "b-qr"
    return (
        f'<article class="bilhete" id="{p["id"]}">'
        f'<header class="b-topo"><p class="rotulo">{esc(p["titulo"])} · <span class="num">{esc(data_txt)}</span></p>'
        f'<p class="b-empresa">{esc(p["empresa"])}</p></header>'
        '<div class="b-horas">'
        f'<div class="b-ponto"><span class="b-hora num">{esc(p["partida"])}</span>'
        f'<span class="b-local">{esc(p["origem"])}</span><span class="b-sub">{esc(p["origem_local"])}</span></div>'
        f'<div class="b-duracao">{esc(p["duracao"])}</div>{chegada}</div>'
        f'<dl class="b-dados">{dados}</dl>'
        '<div class="picote" aria-hidden="true"></div>'
        f'<div class="{qr_classe}"><a class="qr-caixa" href="#qr-{p["id"]}" aria-label="Ampliar QR code">{qr}</a>'
        f'<p class="b-qr-legenda">{esc(p["qr_legenda"])} <a href="#qr-{p["id"]}">Ampliar</a></p></div>'
        f'<p class="b-aviso">{esc(p["aviso"])}</p>'
        f'<div class="acoes">{links}</div>{pagamento}</article>')


def html_tela_qr(p, qr):
    detalhes = " · ".join([p["titulo"], quando(p["data"]), p["partida"]]
                          + [f"{k} {v}" for k, v in p["dados"] if k == "Poltrona"])
    return (
        f'<div class="qr-tela" id="qr-{p["id"]}" role="dialog" aria-label="QR code da {esc(p["titulo"].lower())}">'
        f'<p class="qr-tela-titulo num">{esc(detalhes)}</p>{qr}'
        '<p class="qr-tela-dica">Aumente o brilho da tela.</p>'
        f'<a class="qr-fechar" href="#{p["id"]}">Fechar</a></div>')


def html_contatos(dias):
    itens = []
    for d in dias:
        p = d["pouso"]
        if not p:
            continue
        fones = html_fones(p["fones"], "fones fones-compactos") or '<p class="sem-fone">sem telefone na planilha</p>'
        itens.append(
            f'<li><p class="ag-quando"><a href="#{d["id"]}">{esc(d["dia"])}</a> · {quando(d["data"])} · {esc(p["cidade"])}</p>'
            f'<p class="ag-nome">{esc(p["nome"])}</p>{fones}</li>')
    viagem = "".join(f'<li><p class="ag-nome">{esc(n)}</p>{html_fones([{"numero": t, "href": tel_href(t), "quem": ""}], "fones fones-compactos")}</li>'
                     for n, t in SUPORTE)
    emerg = "".join(f'<li><a class="fone fone-emerg num" href="tel:{t}">{ICONES["fone"]}{t}</a><span>{esc(n)}</span></li>'
                    for n, t in EMERGENCIA)
    return (f'<h3 class="sub-titulo">Pousadas, noite a noite</h3><ul class="agenda">{"".join(itens)}</ul>'
            f'<h3 class="sub-titulo">Viagem</h3><ul class="agenda">{viagem}</ul>'
            f'<h3 class="sub-titulo">Emergência</h3><ul class="emergencia">{emerg}</ul>')


def linkar(texto):
    return re.sub(r"(https?://\S+)", lambda m: f'<a href="{m.group(1)}" target="_blank" rel="noopener">{m.group(1)}</a>', esc(texto))


def html_info(itens, observacao):
    blocos = []
    for k, v in itens:
        if k.lower().startswith("o que levar"):
            lista = "".join(
                f'<li><label><input type="checkbox" id="levar-{i}" data-guardar> <span>{esc(x.strip())}</span></label></li>'
                for i, x in enumerate(re.split(r",\s*(?![^()]*\))", v)))
            blocos.append(f'<div class="info-item info-largo"><dt>{esc(k)}</dt><dd><ul class="levar">{lista}</ul></dd></div>')
        else:
            blocos.append(f'<div class="info-item"><dt>{esc(k)}</dt><dd>{linkar(v)}</dd></div>')
    nota = f'<p class="nota-fonte"><span class="rotulo">Observação da planilha</span>{linkar(observacao)}</p>' if observacao else ""
    return f'<dl class="info">{"".join(blocos)}</dl>{nota}'


def dados_js(dias):
    saida = [{"id": d["id"], "data": d["data"].isoformat(), "dia": d["dia"],
              "km": km_txt(d["km"]) if d["km"] else ""} for d in dias]
    return json.dumps(saida, ensure_ascii=False).replace("</", "<\\/")


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--fragmento", type=Path, help="grava também a versão sem <html>/<head>")
    args = ap.parse_args()

    avisos = []
    wb = openpyxl.load_workbook(PLANILHA)
    titulo, dias, total, observacao = ler_roteiro(wb, avisos)
    info = ler_info(wb)

    qrs = {p["id"]: qr_svg(qr_matriz(p["pdf"]), f"QR code da passagem de {p['titulo'].lower()}") for p in PASSAGENS}
    andando = [d for d in dias if d["km"] > 0]
    noites = sum(1 for d in dias if d["pouso"])
    gerado = dt.date.today()

    trocas = {
        "{{TITULO}}": esc(titulo.split("—")[0].strip()),
        "{{PERIODO}}": f"{dias[0]['data'].day} {MESES[dias[0]['data'].month - 1]} – {dias[-1]['data'].day} {MESES[dias[-1]['data'].month - 1]} {dias[-1]['data'].year}",
        "{{ROTA}}": f"{esc(andando[0]['origem_nome'])} → {esc(andando[-1]['destino_nome'])} · MG",
        "{{TOTAL_KM}}": km_txt(total),
        "{{ETAPAS}}": str(len(andando)),
        "{{NOITES}}": str(noites),
        "{{FAIXA}}": html_faixa(dias),
        "{{DIAS}}": "".join(html_dia(d) for d in dias),
        "{{PASSAGENS}}": "".join(html_passagem(p, qrs[p["id"]]) for p in PASSAGENS),
        "{{TELAS_QR}}": "".join(html_tela_qr(p, qrs[p["id"]]) for p in PASSAGENS),
        "{{CONTATOS}}": html_contatos(dias),
        "{{INFO}}": html_info(info, observacao),
        "{{GERADO}}": f"{gerado.day:02d}/{gerado.month:02d}/{gerado.year}",
        "{{DADOS}}": dados_js(dias),
    }
    modelo = MODELO.read_text(encoding="utf-8")
    for k, v in trocas.items():
        modelo = modelo.replace(k, v)
    sobra = re.search(r"\{\{(?!META\}\})\w+\}\}", modelo)
    if sobra:
        sys.exit("Marcador sem valor: " + sobra.group())
    cabeca, corpo = modelo.split("<!--CORPO-->")

    pagina = (
        '<!doctype html>\n<html lang="pt-BR">\n<head>\n<meta charset="utf-8">\n'
        '<meta name="viewport" content="width=device-width, initial-scale=1, viewport-fit=cover">\n'
        '<meta name="theme-color" content="#1B4E8E">\n'
        '<link rel="manifest" href="manifest.webmanifest">\n'
        '<link rel="icon" href="icon-192.png">\n<link rel="apple-touch-icon" href="apple-touch-icon.png">\n'
        '<meta name="apple-mobile-web-app-capable" content="yes">\n'
        + cabeca.replace("{{META}}", "")
        + "</head>\n<body>" + corpo
        + "<script>if ('serviceWorker' in navigator && location.protocol === 'https:') "
          "navigator.serviceWorker.register('sw.js').catch(function () {});</script>\n</body>\n</html>\n")
    (RAIZ / "index.html").write_text(pagina, encoding="utf-8")

    versao = hashlib.sha1(pagina.encode()).hexdigest()[:10]
    (RAIZ / "sw.js").write_text(SW_MODELO.read_text(encoding="utf-8").replace("{{VERSAO}}", versao), encoding="utf-8")

    if args.fragmento:
        meta = '<meta name="claude:viewport-fit-sides" content="cover">'
        args.fragmento.write_text(cabeca.replace("{{META}}", meta) + corpo, encoding="utf-8")

    print(f"index.html gerado: {len(dias)} dias, {km_txt(total)} km, {noites} noites, {len(pagina) // 1024} KB")
    for a in avisos:
        print("AVISO:", a)


if __name__ == "__main__":
    main()
